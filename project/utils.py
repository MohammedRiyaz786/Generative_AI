import streamlit as st
from langchain.memory import ConversationBufferMemory
from langchain.chains import ConversationalRetrievalChain
from langchain_community.llms import Ollama
from langchain.prompts import PromptTemplate
import requests
import logging
from typing import List
from fastapi import UploadFile
import asyncio
from motor.motor_asyncio import AsyncIOMotorClient
import pickle
import base64
import os
from typing import List, Tuple, BinaryIO
from fastapi import UploadFile
import pdfplumber
import pandas as pd
from langchain_core.documents import Document
from io import BytesIO
from PyPDF2 import PdfReader
from pptx import Presentation
from docx import Document as DocxDocument
from pptx.enum.shapes import MSO_SHAPE_TYPE
import re
import logging
import easyocr
from PIL import Image
import pytesseract
import cv2
import numpy as np
from transformers import TrOCRProcessor, VisionEncoderDecoderModel
import torch
import asyncio

# Initialize OCR components
try:
    reader = easyocr.Reader(['en'])
    trocr_processor = TrOCRProcessor.from_pretrained('microsoft/trocr-base-handwritten')
    trocr_model = VisionEncoderDecoderModel.from_pretrained('microsoft/trocr-base-handwritten')
    device = torch.device('cuda' if torch.cuda.is_available() else 'cpu')
    trocr_model.to(device)
except Exception as e:
    logging.error(f"Error initializing OCR components: {str(e)}")

async def process_image(file: UploadFile) -> Tuple[str, List[Document]]:
    """
    Process image files and extract text using multiple OCR engines.
    Adapted for FastAPI UploadFile.
    """
    try:
        contents = await file.read()
        image = Image.open(BytesIO(contents))
        
        # Convert to numpy array
        image_np = np.array(image)
        
        # Convert to grayscale if needed
        gray = cv2.cvtColor(image_np, cv2.COLOR_RGB2GRAY) if len(image_np.shape) == 3 else image_np
        
        text_results = []
        
        # 1. Pytesseract OCR
        try:
            pytesseract_text = pytesseract.image_to_string(gray)
            if pytesseract_text.strip():
                text_results.append(pytesseract_text)
        except Exception as e:
            logging.warning(f"Pytesseract OCR failed: {str(e)}")
        
        # 2. EasyOCR
        try:
            easy_result = reader.readtext(gray)
            easy_text = " ".join([t[1] for t in easy_result])
            if easy_text.strip():
                text_results.append(easy_text)
        except Exception as e:
            logging.warning(f"EasyOCR failed: {str(e)}")
        
        # 3. TrOCR for handwriting
        try:
            pixel_values = trocr_processor(image, return_tensors="pt").pixel_values.to(device)
            generated_ids = trocr_model.generate(pixel_values)
            trocr_text = trocr_processor.batch_decode(generated_ids, skip_special_tokens=True)[0]
            if trocr_text.strip():
                text_results.append(trocr_text)
        except Exception as e:
            logging.warning(f"TrOCR failed: {str(e)}")
        
        # Extract tables and formulas
        table_text = await extract_tables_from_image(gray)
        formula_text = await extract_formulas_from_image(gray)
        
        if table_text:
            text_results.append(table_text)
        if formula_text:
            text_results.append(formula_text)
        
        combined_text = "\n".join(filter(None, text_results))
        
        if not combined_text.strip():
            raise Exception("No text could be extracted from the image")
        
        return combined_text, [Document(
            page_content=combined_text,
            metadata={'source': 'image', 'filename': file.filename}
        )]
        
    except Exception as e:
        logging.error(f"Error processing image {file.filename}: {str(e)}")
        raise

async def extract_tables_from_image(gray_image: np.ndarray) -> str:
    """Extract tables from images using OpenCV."""
    try:
        denoised = cv2.fastNlMeansDenoising(gray_image)
        thresh = cv2.adaptiveThreshold(denoised, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, 
                                     cv2.THRESH_BINARY_INV, 11, 2)
        
        horizontal_kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (40, 1))
        vertical_kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (1, 40))
        
        horizontal_lines = cv2.morphologyEx(thresh, cv2.MORPH_OPEN, horizontal_kernel)
        vertical_lines = cv2.morphologyEx(thresh, cv2.MORPH_OPEN, vertical_kernel)
        table_mask = cv2.add(horizontal_lines, vertical_lines)
        
        contours, _ = cv2.findContours(table_mask, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
        
        table_text = []
        for contour in contours:
            x, y, w, h = cv2.boundingRect(contour)
            if w > 100 and h > 100:
                roi = gray_image[y:y+h, x:x+w]
                table_content = pytesseract.image_to_string(roi, config='--psm 6')
                
                if not table_content.strip():
                    easy_result = reader.readtext(roi)
                    table_content = " ".join([t[1] for t in easy_result])
                
                if table_content.strip():
                    table_text.append(f"Table Content:\n{table_content}")
        
        return "\n\n".join(table_text) if table_text else ""
        
    except Exception as e:
        logging.warning(f"Table extraction failed: {str(e)}")
        return ""

async def extract_formulas_from_image(gray_image: np.ndarray) -> str:
    """Extract mathematical formulas from images."""
    try:
        _, thresh = cv2.threshold(gray_image, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)
        contours, _ = cv2.findContours(thresh, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
        
        formula_text = []
        for contour in contours:
            x, y, w, h = cv2.boundingRect(contour)
            if w > 30 and h > 30:
                roi = gray_image[y:y+h, x:x+w]
                formula_content = pytesseract.image_to_string(
                    roi,
                    config='--psm 6 -c tessedit_char_whitelist=0123456789+-*/()={}[]abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ'
                )
                
                if formula_content.strip():
                    formula_text.append(f"Formula: {formula_content}")
        
        return "\n".join(formula_text) if formula_text else ""
        
    except Exception as e:
        logging.warning(f"Formula extraction failed: {str(e)}")
        return ""

async def process_pdf(file: UploadFile) -> Tuple[str, List[Document]]:
    """Process PDF files and extract text and tables."""
    try:
        contents = await file.read()
        pdf_file = BytesIO(contents)
        
        text = ""
        documents = []
        
        with pdfplumber.open(pdf_file) as pdf:
            for page_num, page in enumerate(pdf.pages):
                # Extract text
                page_text = page.extract_text(x_tolerance=3, y_tolerance=3) or ""
                text += page_text + "\n"
                
                # Extract tables
                tables = page.extract_tables()
                for table in tables:
                    table_text = "Table:\n"
                    for row in table:
                        filtered_row = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
                        if filtered_row:
                            table_text += " | ".join(filtered_row) + "\n"
                    text += table_text + "\n"
                    
                    documents.append(Document(
                        page_content=table_text,
                        metadata={'source': 'table', 'page': page_num + 1, 'filename': file.filename}
                    ))
                
                documents.append(Document(
                    page_content=page_text,
                    metadata={'source': 'pdf_text', 'page': page_num + 1, 'filename': file.filename}
                ))
        
        return text, documents
        
    except Exception as e:
        logging.error(f"Error processing PDF {file.filename}: {str(e)}")
        raise

async def process_docx(file: UploadFile) -> Tuple[str, List[Document]]:
    """Process Word documents and extract text."""
    try:
        contents = await file.read()
        doc = DocxDocument(BytesIO(contents))
        text = ""
        documents = []
        
        for element in doc.element.body:
            if element.tag.endswith('p'):
                paragraph = element.text
                text += paragraph + "\n"
                
                # Extract formulas
                formulas = re.findall(r'\$.*?\$', paragraph)
                if formulas:
                    text += "Formulas: " + " ".join(formulas) + "\n"
            
            elif element.tag.endswith('tbl'):
                table_text = "Table:\n"
                for row in element.findall('.//w:tr', namespaces=element.nsmap):
                    cells = [cell.text for cell in row.findall('.//w:t', namespaces=element.nsmap)]
                    table_text += " | ".join(cells) + "\n"
                text += table_text + "\n"
        
        documents.append(Document(
            page_content=text,
            metadata={'source': 'word', 'filename': file.filename}
        ))
        
        return text, documents
        
    except Exception as e:
        logging.error(f"Error processing Word document {file.filename}: {str(e)}")
        raise

async def process_pptx(file: UploadFile) -> Tuple[str, List[Document]]:
    """Process PowerPoint files and extract text."""
    try:
        contents = await file.read()
        prs = Presentation(BytesIO(contents))
        text = ""
        documents = []
        
        for slide_num, slide in enumerate(prs.slides):
            slide_text = f"Slide {slide_num + 1}:\n"
            
            for shape in slide.shapes:
                if hasattr(shape, 'text'):
                    shape_text = shape.text.strip()
                    if shape_text:
                        slide_text += shape_text + "\n"
                
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    for subshape in shape.shapes:
                        if hasattr(subshape, 'text'):
                            subshape_text = subshape.text.strip()
                            if subshape_text:
                                slide_text += subshape_text + "\n"
                
                if hasattr(shape, 'has_table') and shape.has_table:
                    table = shape.table
                    table_text = "Table:\n"
                    for row in table.rows:
                        row_text = " | ".join([cell.text.strip() for cell in row.cells])
                        table_text += row_text + "\n"
                    slide_text += table_text + "\n"
            
            formulas = re.findall(r'\$.*?\$', slide_text)
            if formulas:
                slide_text += "Formulas: " + " ".join(formulas) + "\n"
            
            text += slide_text + "\n"
            documents.append(Document(
                page_content=slide_text,
                metadata={'source': f'powerpoint_slide_{slide_num + 1}', 'filename': file.filename}
            ))
        
        return text, documents
        
    except Exception as e:
        logging.error(f"Error processing PowerPoint file {file.filename}: {str(e)}")
        raise

async def process_excel(file: UploadFile) -> Tuple[str, List[Document]]:
    """Process Excel files and extract text."""
    try:
        contents = await file.read()
        excel_data = pd.read_excel(BytesIO(contents), sheet_name=None)
        text = ""
        documents = []
        
        for sheet_name, sheet_data in excel_data.items():
            sheet_text = f"Sheet: {sheet_name}\n"
            
            for _, row in sheet_data.iterrows():
                filtered_row = [str(cell).strip() for cell in row if pd.notna(cell) and str(cell).strip()]
                if filtered_row:
                    sheet_text += " | ".join(filtered_row) + "\n"
            
            text += sheet_text + "\n"
            documents.append(Document(
                page_content=sheet_text,
                metadata={'source': f'excel_sheet_{sheet_name}', 'filename': file.filename}
            ))
        
        return text, documents
        
    except Exception as e:
        logging.error(f"Error processing Excel file {file.filename}: {str(e)}")
        raise

async def process_file(file: UploadFile) -> Tuple[str, List[Document]]:
    """Main function to process any supported file type."""
    try:
        file_extension = os.path.splitext(file.filename)[1].lower()
        
        if file_extension in ['.jpg', '.jpeg', '.png', '.bmp']:
            return await process_image(file)
        elif file_extension == '.pdf':
            return await process_pdf(file)
        elif file_extension in ['.doc', '.docx']:
            return await process_docx(file)
        elif file_extension in ['.ppt', '.pptx']:
            return await process_pptx(file)
        elif file_extension in ['.xls', '.xlsx']:
            return await process_excel(file)
        else:
            raise ValueError(f"Unsupported file type: {file_extension}")
            
    except Exception as e:
        logging.error(f"Error processing file {file.filename}: {str(e)}")
        raise

async def process_files(files: List[UploadFile]) -> Tuple[str, List[Document]]:
    """Process multiple files concurrently."""
    try:
        tasks = [process_file(file) for file in files]
        results = await asyncio.gather(*tasks)
        
        combined_text = "\n\n".join(text for text, _ in results)
        combined_docs = [doc for _, docs in results for doc in docs]
        
        return combined_text, combined_docs
        
    except Exception as e:
        logging.error(f"Error processing files: {str(e)}")
        raise

# Your existing utility functions here (process_image, extract_file_content, etc.)
# Make sure to modify them to work with FastAPI's UploadFile

def init_session_state():
    if 'memory' not in st.session_state:
        st.session_state.memory = ConversationBufferMemory(
            memory_key="chat_history",
            return_messages=True,
            output_key='answer'
        )
    if 'conversation' not in st.session_state:
        st.session_state.conversation = []
    if 'docs_processed' not in st.session_state:
        st.session_state.docs_processed = False

def handle_user_input(user_question: str):
    try:
        # Call API to get response
        response = requests.post(
            'http://localhost:8000/query',
            json={"text": user_question}
        )
        
        if response.status_code == 200:
            answer = response.json()["answer"]
            
            if 'conversation' not in st.session_state:
                st.session_state.conversation = []
                
            st.session_state.conversation.append({
                "user": user_question,
                "assistant": answer
            })
            
            for message in st.session_state.conversation:
                with st.chat_message("user", avatar="🧑"):
                    st.write(message["user"])
                with st.chat_message("assistant", avatar="🤖"):
                    st.write(message["assistant"])
        else:
            st.error("Error getting response from API")
            
    except Exception as e:
        logging.error(f"Error in handle_user_input: {str(e)}", exc_info=True)
        st.error(f"An error occurred: {str(e)}")

def create_qa_chain(vector_store):
    prompt_template = """You are a polite, respectful, and efficient AI assistant.

    IF the user's message matches ANY of these patterns:
    - "Hi", "Hello", "Hey", "Hii", "Hola" (just greeting)
    - "My name is [any name]" 
    - "I am [any name]"
    - "[any greeting] my name is [any name]"
    - "[any greeting] I am [any name]"
    THEN respond only with: "Hello! How can I assist you today?"

    OTHERWISE:
    1. Use only the provided information:
    - Context: {context}
    - Chat History: {chat_history} 
    - Current Question: {question}

    2. Your response must be:
    - Direct and to-the-point
    - Based only on given context and history
    - Without any explanations about your capabilities
    - Without mentioning sources or references

    3. If the answer cannot be found in context or history:
    Response should be only: "I apologize, but I don't have enough information to answer your question."

    4. Never start responses with:
    - "Based on..."
    - "According to..." 
    - "I understand..."
    - "Let me..."

    5. Never end responses with:
    - "Is there anything else..."
    - "Let me know if..."
    - "Feel free to..."

    6. If the user uses any abusive or inappropriate language, respond politely and avoid escalation:
    "I apologize, but I don't engage with that type of language. How else can I assist you today?"

    Question: {question}"""

    PROMPT = PromptTemplate(
        template=prompt_template,
        input_variables=["context", "question", "chat_history"]
    )

    llm = Ollama(model="llama3.1", temperature=0.1)
    
    retriever = vector_store.as_retriever(
        search_type="mmr",
        search_kwargs={"k": 5, "fetch_k": 20}
    )

    qa_chain = ConversationalRetrievalChain.from_llm(
        llm=llm,
        retriever=retriever,
        memory=st.session_state.memory,
        combine_docs_chain_kwargs={"prompt": PROMPT},
        return_source_documents=True
    )
    
    return qa_chain